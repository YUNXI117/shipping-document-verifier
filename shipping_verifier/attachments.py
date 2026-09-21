"""Read attachment content through one format-independent interface."""

from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


class AttachmentError(Exception):
    """Base exception for attachment-reading failures."""


class UnsupportedAttachmentType(AttachmentError):
    """Raised when no reader is available for a file extension."""


def resolve_attachment(data_dir: Path, attachment: str) -> Path:
    """Resolve an attachment path and prevent access outside the dataset."""
    root = data_dir.resolve()
    path = (root / attachment).resolve()
    if not path.is_relative_to(root):
        raise AttachmentError(f"Attachment is outside the dataset: {attachment}")
    if not path.is_file():
        raise AttachmentError(f"Attachment does not exist: {attachment}")
    return path


def read_attachment(data_dir: Path, attachment: str) -> str:
    """Extract text from a supported attachment."""
    path = resolve_attachment(data_dir, attachment)
    readers = {
        ".txt": read_txt,
        ".pdf": read_pdf,
        ".docx": read_docx,
        ".xlsx": read_xlsx,
    }
    reader = readers.get(path.suffix.lower())
    if reader is not None:
        try:
            text = reader(path)
        except AttachmentError:
            raise
        except Exception as error:
            raise AttachmentError(f"Could not read attachment: {attachment}") from error
        if not text.strip():
            raise AttachmentError(f"Attachment contains no readable text: {attachment}")
        return text
    raise UnsupportedAttachmentType(
        f"Unsupported attachment type: {path.suffix.lower() or '<none>'}"
    )


def read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def read_pdf(path: Path) -> str:
    pages = (page.extract_text() or "" for page in PdfReader(path).pages)
    return "\n".join(pages)


def read_docx(path: Path) -> str:
    document = Document(path)
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                lines.append(" | ".join(values))
    return "\n".join(lines)


def read_xlsx(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        lines: list[str] = []
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None]
                if values:
                    lines.append(" | ".join(values))
        return "\n".join(lines)
    finally:
        workbook.close()
